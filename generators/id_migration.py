from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

from generators.id_analysis import analyze_id_column


def migrate_ids_to_static(
    xlsx_path: str | Path,
    output_path: str | Path,
    sheet_name: str | None = None,
) -> dict:
    """Crea una copia XLSX con gli ID formula trasformati in valori statici."""

    xlsx_path = Path(xlsx_path)
    output_path = Path(output_path)

    if xlsx_path.resolve() == output_path.resolve():
        raise ValueError("Il file di output deve essere diverso dall'originale.")

    if output_path.exists():
        raise FileExistsError(f"Il file di output esiste già: {output_path}")

    report = analyze_id_column(xlsx_path, sheet_name=sheet_name)

    if report["status"] not in ("valid", "formula_based"):
        raise ValueError(f"Analisi ID non valida: {report}")

    wb = load_workbook(xlsx_path, data_only=False)
    values_wb = load_workbook(xlsx_path, data_only=True)

    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        values_ws = values_wb[ws.title]

        id_col = report["idColumn"]
        converted = []

        for row in report["formulaRows"]:
            value = values_ws.cell(row, id_col).value
            ws.cell(row, id_col).value = int(value)
            converted.append(row)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)

        return {
            "source": str(xlsx_path),
            "output": str(output_path),
            "convertedRows": converted,
            "convertedCount": len(converted),
            "dataRows": report["dataRows"],
        }

    finally:
        wb.close()
        values_wb.close()
