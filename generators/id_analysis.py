from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any, Dict

from openpyxl import load_workbook


def analyze_id_column(
    xlsx_path: str | Path,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> Dict[str, Any]:
    """Analizza gli ID senza modificare il workbook."""

    wb = load_workbook(xlsx_path, data_only=False, read_only=False)
    values_wb = load_workbook(xlsx_path, data_only=True, read_only=False)

    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        values_ws = values_wb[ws.title]

        headers = [
            str(cell.value or "").strip().lower()
            for cell in ws[header_row]
        ]

        id_columns = [
            i + 1 for i, name in enumerate(headers)
            if name == "id"
        ]

        # Consideriamo soltanto le colonne che hanno un'intestazione.
        data_columns = [
            i + 1 for i, name in enumerate(headers)
            if name
        ]

        occupied_rows = []

        for row in range(header_row + 1, ws.max_row + 1):
            values = [
                values_ws.cell(row, col).value
                for col in data_columns
            ]

            if any(
                value is not None and str(value).strip() != ""
                for value in values
            ):
                occupied_rows.append(row)

        if not id_columns:
            return {
                "status": "missing",
                "idColumn": None,
                "dataRows": len(occupied_rows),
                "occupiedRows": occupied_rows,
            }

        if len(id_columns) > 1:
            return {
                "status": "invalid",
                "reason": "multiple_id_columns",
                "idColumns": id_columns,
            }

        col = id_columns[0]

        ids = []
        missing_rows = []
        invalid_rows = []
        formula_rows = []
        uncached_formula_rows = []

        for row in occupied_rows:
            raw = ws.cell(row, col).value
            value = values_ws.cell(row, col).value

            if isinstance(raw, str) and raw.startswith("="):
                formula_rows.append(row)
                if value is None:
                    uncached_formula_rows.append(row)

            if value is None or str(value).strip() == "":
                missing_rows.append(row)
                continue

            if isinstance(value, bool):
                invalid_rows.append(row)
                continue

            try:
                number = int(value)

                if str(value).strip() not in (
                    str(number),
                    f"{number}.0",
                ):
                    raise ValueError

                if number < 1:
                    raise ValueError

                ids.append(number)

            except (ValueError, TypeError):
                invalid_rows.append(row)

        counts = Counter(ids)

        duplicate_ids = sorted(
            number
            for number, count in counts.items()
            if count > 1
        )

        status = "valid"

        if missing_rows or invalid_rows or duplicate_ids:
            status = "invalid"
        elif formula_rows:
            status = "formula_based"

        return {
            "status": status,
            "idColumn": col,
            "dataRows": len(occupied_rows),
            "occupiedRows": occupied_rows,
            "missingRows": missing_rows,
            "invalidRows": invalid_rows,
            "duplicateIds": duplicate_ids,
            "formulaRows": formula_rows,
            "uncachedFormulaRows": uncached_formula_rows,
        }

    finally:
        wb.close()
        values_wb.close()
