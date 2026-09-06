from __future__ import annotations

import json
import re
import uuid
import zipfile
from pathlib import Path
from typing import Any, Dict

import requests

from flask import (
    Flask,
    jsonify,
    redirect,
    request,
    send_file,
)
from werkzeug.utils import secure_filename

from pipeline.build_schema_from_input import build_schema_from_directory
from app.builder_state import build_builder_state
from generators.builder_generators import generate_all


ROOT = Path(__file__).resolve().parent
DOCS_DIR = ROOT / "docs"
WORKDIR = ROOT / "workdir" / "flask_sessions"
DATA_DIR = ROOT / "data"
APPS_REGISTRY_PATH = DATA_DIR / "apps_registry.json"

# =========================================================
# APPS CONFIG (slug -> hidden backend config)
# =========================================================


PUBLIC_MODES = {"meta", "schema", "view"}
PROTECTED_MODES = {"insert", "getById", "update", "delete"}
ALL_PROXY_MODES = PUBLIC_MODES | PROTECTED_MODES

# =========================================================
# FLASK APP
# =========================================================
app = Flask(
    __name__,
    static_folder=str(DOCS_DIR),
    static_url_path="",
)

WORKDIR.mkdir(parents=True, exist_ok=True)

def slugify_app_name(value: str) -> str:
    value = str(value or "").strip()
    value = re.sub(r"([a-z0-9])([A-Z])", r"\1-\2", value)
    value = value.lower()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    value = value.strip("-")
    return value or "generated-app"

def load_apps_registry() -> Dict[str, Dict[str, str]]:
    if not APPS_REGISTRY_PATH.exists():
        return {}

    raw = APPS_REGISTRY_PATH.read_text(encoding="utf-8").strip()

    if not raw:
        return {}

    return json.loads(raw)


def save_apps_registry(registry: Dict[str, Dict[str, str]]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    APPS_REGISTRY_PATH.write_text(
        json.dumps(registry, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def register_app_config(slug: str, web_app_url: str, api_key: str) -> Dict[str, str]:
    safe_slug = slugify_app_name(slug)

    registry = load_apps_registry()
    registry[safe_slug] = {
        "web_app_url": web_app_url.strip(),
        "api_key": api_key.strip(),
    }

    save_apps_registry(registry)
    return registry[safe_slug]


# =========================================================
# HELPERS
# =========================================================
def write_json(path: Path, data: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def builder_state_to_runtime_data(
    builder_state: Dict[str, Any]
) -> tuple[Dict[str, Any], Dict[str, Any]]:
    project = dict(builder_state.get("project", {}) or {})
    fields = list(builder_state.get("fields", []) or [])
    constraints = dict(builder_state.get("constraints", {}) or {})
    enums = dict(builder_state.get("enums", {}) or {})
    enum_styles = dict(builder_state.get("enumStyles", {}) or {})

    headers = [f["name"] for f in fields]

    required_on_insert = [
        f["name"]
        for f in fields
        if bool(f.get("required")) and not bool(f.get("computed"))
    ]

    optional_on_insert = [
        f["name"]
        for f in fields
        if not bool(f.get("required")) and not bool(f.get("computed"))
    ]

    computed = [
        f["name"]
        for f in fields
        if bool(f.get("computed"))
    ]

    visible_in_form = [
        f["name"]
        for f in fields
        if bool(f.get("visibleInForm"))
    ]

    visible_in_viewer = [
        f["name"]
        for f in fields
        if bool(f.get("visibleInViewer"))
    ]

    fields_schema = {
        "headers": headers,
        "requiredOnInsert": required_on_insert,
        "optionalOnInsert": optional_on_insert,
        "computed": computed,
        "visibleInForm": visible_in_form,
        "visibleInViewer": visible_in_viewer,
        "constraints": constraints,
        "enums": enums,
        "enumStyles": enum_styles,
        "fields": fields,
    }

    return project, fields_schema


def create_session_dir() -> Path:
    session_id = uuid.uuid4().hex[:12]
    session_dir = WORKDIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)
    return session_dir


def get_session_dir(session_id: str) -> Path:
    safe = secure_filename(session_id)
    path = WORKDIR / safe

    if not path.exists() or not path.is_dir():
        raise FileNotFoundError(f"Session not found: {session_id}")

    return path


def process_builder_input(
    session_dir: Path,
    input_dir: Path,
    sheet_name: str,
    xlsx_sheet_name: str | None = None,
) -> tuple[Dict[str, Any], Path, Path]:
    parsed_schema_path = session_dir / "parsed.schema.json"
    builder_state_path = session_dir / "builder_state.json"

    effective_xlsx_sheet_name = xlsx_sheet_name or sheet_name

    parsed_schema = build_schema_from_directory(
        input_dir=input_dir,
        output_json_path=parsed_schema_path,
        enum_field_name="rating",
        sheet_name=effective_xlsx_sheet_name,
        debug=True,
    )

    builder_state = build_builder_state(parsed_schema)

    project_slug = slugify_app_name(sheet_name)

    builder_state["project"]["sheetName"] = sheet_name
    builder_state["project"]["projectName"] = sheet_name
    builder_state["project"]["projectSlug"] = project_slug
    builder_state["project"]["backendName"] = f"{project_slug}-backend"

    write_json(builder_state_path, builder_state)

    return builder_state, parsed_schema_path, builder_state_path


# =========================================================
# PROXY HELPERS
# =========================================================
def get_app_config(slug: str) -> Dict[str, str]:
    safe_slug = slugify_app_name(slug)
    registry = load_apps_registry()
    config = registry.get(safe_slug)

    if not config:
        raise FileNotFoundError(f"Unknown app slug: {slug}")

    return config


def parse_apps_script_response(text: str) -> Dict[str, Any]:
    payload = text.strip()

    # JSONP style: callback({...});
    if "(" in payload and payload.endswith(");"):
        start = payload.find("(")
        payload = payload[start + 1:-2]

    return json.loads(payload)


def call_apps_script_proxy(
    slug: str,
    mode: str,
    params: Dict[str, Any],
) -> Dict[str, Any]:
    if mode not in ALL_PROXY_MODES:
        raise ValueError(f"Invalid mode: {mode}")

    config = get_app_config(slug)

    outbound = {
    "mode": mode,
    "cb": "djungoProxy",
}

    for key, value in params.items():
        if key in {"mode", "apiKey", "webApp", "cb", "_"}:
            continue
        outbound[key] = value

    if mode in PROTECTED_MODES:
        outbound["apiKey"] = config["api_key"]

    resp = requests.get(
        config["web_app_url"],
        params=outbound,
        timeout=45,
    )

    resp.raise_for_status()

    return parse_apps_script_response(resp.text)


# =========================================================
# ROUTES
# =========================================================
@app.get("/")
def home():
    return redirect("/generate.html")


@app.get("/health")
def health():
    return jsonify({
        "ok": True,
        "service": "sheet-builder-flask",
    })


# ---------------------------------------------------------
# PARSE INPUT FILES
# ---------------------------------------------------------
@app.post("/api/parse")
def api_parse():
    try:
        xlsx_file = request.files.get("xlsx")
        html_file = request.files.get("html")
        sheet_name = (request.form.get("sheet_name") or "").strip()

        if not xlsx_file:
            return jsonify({"ok": False, "error": "Missing xlsx file"}), 400

        if not html_file:
            return jsonify({"ok": False, "error": "Missing html file"}), 400

        if not sheet_name:
            return jsonify({"ok": False, "error": "Missing sheet_name"}), 400

        session_dir = create_session_dir()

        input_dir = session_dir / "input"
        input_dir.mkdir(parents=True, exist_ok=True)

        xlsx_name = secure_filename(xlsx_file.filename or "input.xlsx")
        html_name = secure_filename(html_file.filename or "input.html")

        xlsx_path = input_dir / xlsx_name
        html_path = input_dir / html_name

        xlsx_file.save(xlsx_path)
        html_file.save(html_path)

        builder_state, parsed_schema_path, builder_state_path = process_builder_input(
            session_dir=session_dir,
            input_dir=input_dir,
            sheet_name=sheet_name,
        )

        return jsonify({
            "ok": True,
            "session_id": session_dir.name,
            "builder_state": builder_state,
            "paths": {
                "parsed_schema": str(parsed_schema_path),
                "builder_state": str(builder_state_path),
            },
        })

    except Exception as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 500


# ---------------------------------------------------------
# IMPORT FROM GOOGLE SHEET
# ---------------------------------------------------------
@app.post("/api/google-sheet/import")
def api_google_sheet_import():
    try:
        xlsx_file = request.files.get("xlsx")
        zip_file = request.files.get("web_export")

        spreadsheet_id = (request.form.get("spreadsheet_id") or "").strip()
        spreadsheet_name = (request.form.get("spreadsheet_name") or "").strip()
        sheet_id = (request.form.get("sheet_id") or "").strip()
        sheet_name = (request.form.get("sheet_name") or "").strip()

        if not xlsx_file:
            return jsonify({"ok": False, "error": "Missing xlsx file"}), 400

        if not zip_file:
            return jsonify({"ok": False, "error": "Missing web_export file"}), 400

        if not spreadsheet_id:
            return jsonify({"ok": False, "error": "Missing spreadsheet_id"}), 400

        if not sheet_id:
            return jsonify({"ok": False, "error": "Missing sheet_id"}), 400

        if not sheet_name:
            return jsonify({"ok": False, "error": "Missing sheet_name"}), 400

        session_dir = create_session_dir()

        input_dir = session_dir / "input"
        input_dir.mkdir(parents=True, exist_ok=True)

        xlsx_path = input_dir / "input.xlsx"
        zip_path = session_dir / "google-web-export.zip"

        xlsx_file.save(xlsx_path)
        zip_file.save(zip_path)

        extract_dir = session_dir / "web_export"
        extract_dir.mkdir(parents=True, exist_ok=True)

        with zipfile.ZipFile(zip_path, "r") as archive:
            for member in archive.infolist():
                member_path = Path(member.filename)

                if member_path.is_absolute() or ".." in member_path.parts:
                    raise ValueError(
                        f"Unsafe ZIP member path: {member.filename}"
                    )

                target_path = extract_dir / member_path

                if member.is_dir():
                    target_path.mkdir(parents=True, exist_ok=True)
                    continue

                target_path.parent.mkdir(parents=True, exist_ok=True)

                with archive.open(member, "r") as source:
                    with target_path.open("wb") as target:
                        target.write(source.read())

        html_candidates = [
            path
            for path in extract_dir.rglob("*.html")
            if path.stem == sheet_name
        ]

        if len(html_candidates) != 1:
            raise ValueError(
                "Could not uniquely identify HTML for selected sheet "
                f"{sheet_name!r}; found {len(html_candidates)} matches"
            )

        selected_html = html_candidates[0]
        html_path = input_dir / "input.html"
        html_path.write_bytes(selected_html.read_bytes())

        from openpyxl import load_workbook

        wb = load_workbook(
            filename=xlsx_path,
            read_only=True,
            data_only=False,
        )

        try:
            if sheet_name in wb.sheetnames:
                xlsx_sheet_name = sheet_name
            elif sheet_name[:31] in wb.sheetnames:
                xlsx_sheet_name = sheet_name[:31]
            else:
                raise ValueError(
                    "Could not identify selected sheet inside XLSX. "
                    f"Google sheet name: {sheet_name!r}; "
                    f"XLSX worksheets: {wb.sheetnames!r}"
                )
        finally:
            wb.close()

        builder_state, parsed_schema_path, builder_state_path = process_builder_input(
            session_dir=session_dir,
            input_dir=input_dir,
            sheet_name=sheet_name,
            xlsx_sheet_name=xlsx_sheet_name,
        )

        builder_state["project"]["source"] = {
            "type": "google_sheet",
            "spreadsheetId": spreadsheet_id,
            "spreadsheetName": spreadsheet_name,
            "sheetId": sheet_id,
            "sheetName": sheet_name,
        }

        write_json(builder_state_path, builder_state)

        return jsonify({
            "ok": True,
            "session_id": session_dir.name,
            "builder_state": builder_state,
            "google_source": {
                "spreadsheet_id": spreadsheet_id,
                "spreadsheet_name": spreadsheet_name,
                "sheet_id": sheet_id,
                "sheet_name": sheet_name,
            },
            "paths": {
                "parsed_schema": str(parsed_schema_path),
                "builder_state": str(builder_state_path),
            },
        })

    except zipfile.BadZipFile:
        return jsonify({
            "ok": False,
            "error": "Invalid Google web export ZIP",
        }), 400

    except Exception as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 500


# ---------------------------------------------------------
# GENERATE FINAL FILES
# ---------------------------------------------------------

@app.get("/api/session/<session_id>")
def api_get_session(session_id: str):
    try:
        safe_session_id = secure_filename(session_id)

        if not safe_session_id or safe_session_id != session_id:
            return jsonify({
                "ok": False,
                "error": "Invalid session_id",
            }), 400

        session_dir = get_session_dir(session_id)
        builder_state_path = session_dir / "builder_state.json"

        if not builder_state_path.exists():
            return jsonify({
                "ok": False,
                "error": "Builder state not found",
            }), 404

        builder_state = json.loads(
            builder_state_path.read_text(encoding="utf-8")
        )

        return jsonify({
            "ok": True,
            "session_id": session_id,
            "builder_state": builder_state,
        })

    except FileNotFoundError as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 404

    except Exception as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 500


@app.post("/api/generate")
def api_generate():
    try:
        payload = request.get_json(silent=True) or {}

        session_id = (payload.get("session_id") or "").strip()
        builder_state = payload.get("builder_state")

        if not session_id:
            return jsonify({
                "ok": False,
                "error": "Missing session_id",
            }), 400

        if not isinstance(builder_state, dict):
            return jsonify({
                "ok": False,
                "error": "Missing or invalid builder_state",
            }), 400

        session_dir = get_session_dir(session_id)

        approved_path = session_dir / "approved_configuration.json"
        write_json(approved_path, builder_state)

        project_config, fields_schema = builder_state_to_runtime_data(
            builder_state
        )

        generated = generate_all(project_config, fields_schema)

        generated_root = session_dir / "generated"
        generated_root.mkdir(parents=True, exist_ok=True)

        for relative_path, content in generated.items():
            target = generated_root / relative_path
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding="utf-8")
        
        # Also publish generated runtime files to the Flask static folder.
        # This makes the updated CRUD/viewer immediately visible in the browser.
        for relative_path, content in generated.items():
            target = ROOT / relative_path
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding="utf-8")

        return jsonify({
            "ok": True,
            "session_id": session_id,
            "downloads": {
                "index_html": f"/api/download/{session_id}/index.html",
                "viewer_html": f"/api/download/{session_id}/viewer.html",
                "codice_gs": f"/api/download/{session_id}/codice.gs",
            },
            "generated_root": str(generated_root),
        })

    except Exception as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 500


# ---------------------------------------------------------
# DOWNLOAD GENERATED FILES
# ---------------------------------------------------------
@app.get("/api/download/<session_id>/<filename>")
def api_download(session_id: str, filename: str):
    try:
        session_dir = get_session_dir(session_id)
        generated_root = session_dir / "generated"

        mapping = {
            "index.html": generated_root / "docs" / "index.html",
            "viewer.html": generated_root / "docs" / "viewer.html",
            "codice.gs": generated_root / "codice.gs",
        }

        if filename not in mapping:
            return jsonify({
                "ok": False,
                "error": "Invalid filename",
            }), 400

        file_path = mapping[filename]

        if not file_path.exists():
            return jsonify({
                "ok": False,
                "error": "Generated file not found",
            }), 404

        return send_file(
            file_path,
            as_attachment=True,
            download_name=filename,
        )

    except Exception as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 500

@app.get("/run/<session_id>/<filename>")
def run_generated_app(session_id: str, filename: str):
    try:
        session_dir = get_session_dir(session_id)
        generated_root = session_dir / "generated"

        mapping = {
            "index.html": generated_root / "docs" / "index.html",
            "viewer.html": generated_root / "docs" / "viewer.html",
        }

        if filename not in mapping:
            return jsonify({"ok": False, "error": "Invalid filename"}), 400

        file_path = mapping[filename]
        if not file_path.exists():
            return jsonify({"ok": False, "error": "Generated file not found"}), 404

        return send_file(file_path, as_attachment=False)

    except Exception as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 500
        
        
# ---------------------------------------------------------
# NEW PROXY ROUTE (SAFE CRUD)
# ---------------------------------------------------------
@app.post("/api/apps/register")
def api_register_app():
    try:
        payload = request.get_json(silent=True) or {}

        slug = (payload.get("slug") or "").strip()
        web_app_url = (payload.get("web_app_url") or "").strip()
        api_key = (payload.get("api_key") or "").strip()

        if not slug:
            return jsonify({"ok": False, "error": "Missing slug"}), 400

        if not web_app_url:
            return jsonify({"ok": False, "error": "Missing web_app_url"}), 400

        if not api_key:
            return jsonify({"ok": False, "error": "Missing api_key"}), 400

        safe_slug = slugify_app_name(slug)
        config = register_app_config(
            slug=safe_slug,
            web_app_url=web_app_url,
            api_key=api_key,
        )

        return jsonify({
            "ok": True,
            "slug": safe_slug,
            "registered": True,
            "web_app_url_present": bool(config.get("web_app_url")),
            "api_key_present": bool(config.get("api_key")),
        })

    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500
    
@app.route("/api/apps/<slug>/<mode>", methods=["GET", "POST"])
def api_app_proxy(slug: str, mode: str):
    try:
        params: Dict[str, Any] = {}

        if request.method == "GET":
            params.update(request.args.to_dict())

        else:
            if request.is_json:
                params.update(request.get_json(silent=True) or {})
            else:
                params.update(request.form.to_dict())

        data = call_apps_script_proxy(
            slug=slug,
            mode=mode,
            params=params,
        )

        return jsonify(data)

    except FileNotFoundError as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 404

    except ValueError as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 400

    except requests.RequestException as exc:
        return jsonify({
            "ok": False,
            "error": f"Apps Script request failed: {exc}",
        }), 502

    except Exception as exc:
        return jsonify({
            "ok": False,
            "error": str(exc),
        }), 500


# =========================================================
# MAIN
# =========================================================
if __name__ == "__main__":
    app.run(
        debug=True,
        host="127.0.0.1",
        port=5000,
    )