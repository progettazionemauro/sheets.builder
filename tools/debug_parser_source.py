from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from generators.import_from_sheet import import_schema_from_xlsx


SESSIONS_DIR = ROOT / "workdir" / "flask_sessions"


def latest_file(pattern: str) -> Path:
    files = sorted(
        SESSIONS_DIR.glob(pattern),
        key=lambda p: p.stat().st_mtime,
    )

    if not files:
        raise SystemExit(f"No files found for pattern: {pattern}")

    return files[-1]


def print_section(title: str) -> None:
    print("\n" + "=" * 80)
    print(title)
    print("=" * 80)


def main() -> None:
    xlsx_path = latest_file("*/input/*.xlsx")
    html_path = latest_file("*/input/*.html")
    builder_state_path = latest_file("*/builder_state.json")

    print_section("LATEST INPUT FILES")
    print("XLSX:", xlsx_path)
    print("HTML:", html_path)
    print("BUILDER STATE:", builder_state_path)

    print_section("XLSX SHEET NAMES")
    wb = load_workbook(xlsx_path, read_only=True, data_only=False)

    for name in wb.sheetnames:
        print("-", repr(name))

    print_section("HEADERS READ DIRECTLY FROM XLSX SHEET = ISI2013")

    try:
        schema = import_schema_from_xlsx(
            xlsx_path=xlsx_path,
            sheet_name="ISI2013",
            header_row=1,
            sample_row=2,
            debug=True,
        )

        headers = schema.get("headers", [])

        print("Headers count:", len(headers))

        for h in headers:
            print("-", h)

    except Exception as exc:
        print("ERROR reading sheet ISI2013:", repr(exc))

    print_section("SEARCH OLD/NEW MARKERS IN LATEST HTML")

    html_text = html_path.read_text(
        encoding="utf-8",
        errors="replace",
    )

    markers = [
        "Agri-food",
        "Garanzia Giovani",
        "iNIDIRIZZO",
        "Sviluppo Lazio",
        "ISI2013",
        "Isi_2013",
    ]

    for marker in markers:
        result = "FOUND" if marker in html_text else "not found"
        print(f"{marker!r}: {result}")

    print_section("SEARCH OLD/NEW MARKERS IN LATEST BUILDER STATE")

    state_text = builder_state_path.read_text(
        encoding="utf-8",
        errors="replace",
    )

    for marker in markers:
        result = "FOUND" if marker in state_text else "not found"
        print(f"{marker!r}: {result}")


if __name__ == "__main__":
    main()