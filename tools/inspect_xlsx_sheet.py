from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

xlsx = sorted(
    Path("workdir/flask_sessions").glob("*/input/*.xlsx"),
    key=lambda p: p.stat().st_mtime
)[-1]

wb = load_workbook(xlsx, data_only=False)
ws = wb["ISI2013"]

print("XLSX:", xlsx)
print("SHEET:", ws.title)
print()

for col in range(1, ws.max_column + 1):
    letter = get_column_letter(col)
    value = ws.cell(row=1, column=col).value
    hidden = ws.column_dimensions[letter].hidden
    print(f"{letter}: {value!r} | hidden={hidden}")